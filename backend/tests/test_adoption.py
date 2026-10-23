"""采纳与反馈：第 10 项要求「记录用户是否采纳建议」，并保留未知。

关键行为：没有反馈时默认 unknown（显示「未知」），而不是 assumed「未采纳」——把沉默当成拒绝
会污染后续的采纳率与推荐评估。
"""
import io
import json

import pytest
from openpyxl import load_workbook

from backend import tenants, tracking, weekly_report
from backend.database import connect, initialize


@pytest.fixture
def db(tmp_path):
    path = tmp_path / "main.sqlite3"
    initialize(path)
    with connect(path) as con:
        con.execute("INSERT INTO sources VALUES(?,?,?,?,?,?)",
                    ("s1", "来源", "government", 0, "https://a.test", "2026-10-01T00:00:00+00:00"))
        con.execute("INSERT INTO districts(id,payload) VALUES(?,?)",
                    ("suzhou", json.dumps({"id": "suzhou", "name": "苏州市", "kind": "city"})))
        try:
            con.execute("ALTER TABLE events ADD COLUMN canonical_id TEXT")
        except Exception:
            pass
    return path


@pytest.fixture
def base(tmp_path):
    directory = tmp_path / "data"
    directory.mkdir()
    return directory


def test_adoption_defaults_to_unknown(db):
    tracking.set_followup(db, tenant_id="hotel-a", event_key="e1", state="todo")
    row = tracking.list_followups(db, "hotel-a")[0]
    assert row["adopted"] == "unknown"
    assert row["adoptedLabel"] == "未知"


def test_adoption_persists_with_feedback(db):
    tracking.set_followup(db, tenant_id="hotel-a", event_key="e1", state="done",
                          adopted="adopted", feedback="已按建议联系主办方")
    row = tracking.list_followups(db, "hotel-a")[0]
    assert row["adopted"] == "adopted"
    assert row["adoptedLabel"] == "已采纳"
    assert row["feedback"] == "已按建议联系主办方"


def test_adoption_is_kept_when_only_state_changes(db):
    """之后只改状态不应把已有的采纳结论清空。"""
    tracking.set_followup(db, tenant_id="hotel-a", event_key="e1", state="todo",
                          adopted="adopted")
    tracking.set_followup(db, tenant_id="hotel-a", event_key="e1", state="done")
    row = tracking.list_followups(db, "hotel-a")[0]
    assert row["state"] == "done"
    assert row["adopted"] == "adopted"


def test_rejects_unknown_adoption_value(db):
    with pytest.raises(ValueError, match="采纳状态"):
        tracking.set_followup(db, tenant_id="hotel-a", event_key="e1", state="todo",
                              adopted="maybe")


def test_older_database_gets_the_adoption_columns_added(db):
    """迁移前建的库没有 adopted/feedback 列，ensure_schema 要能补上。"""
    with connect(db) as con:
        con.execute("DROP TABLE IF EXISTS event_followups")
        con.execute(
            "CREATE TABLE event_followups (tenant_id TEXT NOT NULL, event_key TEXT NOT NULL, "
            "state TEXT NOT NULL, note TEXT, updated_at TEXT NOT NULL, updated_by TEXT, "
            "PRIMARY KEY (tenant_id, event_key))")
    tracking.set_followup(db, tenant_id="hotel-a", event_key="e1", state="todo",
                          adopted="declined")
    assert tracking.list_followups(db, "hotel-a")[0]["adoptedLabel"] == "未采纳"


def test_weekly_report_shows_adoption_or_unknown(db, base):
    tenants.create_tenant(db, tenant_id="hotel-a", name="A 酒店")
    tenants.add_member(db, tenant_id="hotel-a", login="owner-a", role="owner")
    tracking.set_followup(db, tenant_id="hotel-a", event_key="e1", state="done",
                          adopted="adopted", feedback="已联系")
    tracking.set_followup(db, tenant_id="hotel-a", event_key="e2", state="todo")

    report = weekly_report.collect(db, base, "hotel-a")
    payload = weekly_report.build_xlsx(report)
    book = load_workbook(io.BytesIO(payload))
    sheet = book["待办"]
    text = " ".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value)
    assert "是否采纳" in text
    assert "已采纳" in text
    assert "未知" in text
    assert "已联系" in text
