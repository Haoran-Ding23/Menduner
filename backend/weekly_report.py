"""Weekly report for one hotel: its own numbers, and the public events it should plan around.

## What this is, and what it deliberately is not

The report states **what was measured**: the hotel's own uploaded figures for the week, and the
public events that overlap that week. It does not forecast demand, recommend a rate, or estimate
revenue. Those are excluded by the project's rules, and the report says so in its footer rather
than leaving the reader to wonder why the numbers are plain.

## Scope, and how it is established

Numbers come only from the tenant's own private database - the report is generated for the
tenant resolved from the session, and there is no parameter that could select another hotel.

Events are public. Where they are listed relative to the hotel depends on whether the hotel's
location is known:

  * the tenant is tied to a row in ``hotels`` when its name matches one, and then events are
    filtered to 10 km - the same radius the map's demand layer uses;
  * when it is not tied, the report lists the week's events across the city and says so. It does
    not invent a location to make the section look more specific than it is.

## Period

The last complete Monday-Sunday week. A report covering a partial week would mix days that
already happened with days that have not, and any average over that is not a weekly figure.
"""
from __future__ import annotations

import io
import math
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from backend import tenants
from backend.database import connect

# The radius is the same one the map draws, so a hotel reading both sees one number.
NEARBY_KM = 10.0

METRIC_LABELS = {
    "occupancy_pct": ("入住率", "%"),
    "adr": ("平均房价", ""),
    "room_nights": ("间夜", ""),
    "channel_share": ("渠道占比", "%"),
    "room_type_share": ("房型结构", "%"),
}

WEEKDAYS = ("周一", "周二", "周三", "周四", "周五", "周六", "周日")


def week_range(today: date | None = None) -> tuple[date, date]:
    """The last complete Monday-Sunday week, ending before today's week."""
    today = today or datetime.now(timezone.utc).date()
    this_monday = today - timedelta(days=today.weekday())
    end = this_monday - timedelta(days=1)
    start = end - timedelta(days=6)
    return start, end


def _haversine_km(a: tuple[float, float], b: tuple[float, float]) -> float:
    radius = 6371.0
    la1, lo1, la2, lo2 = map(math.radians, (a[1], a[0], b[1], b[0]))
    dla, dlo = la2 - la1, lo2 - lo1
    h = math.sin(dla / 2) ** 2 + math.cos(la1) * math.cos(la2) * math.sin(dlo / 2) ** 2
    return 2 * radius * math.asin(math.sqrt(h))


def _tenant_hotel(db_path: Path, tenant: str) -> dict | None:
    """The hotel record this tenant names, when the names match exactly.

    A fuzzy match would silently attach a report to the wrong building, so only an exact name
    is accepted; otherwise the caller is told the location is unknown.
    """
    with connect(db_path) as con:
        row = con.execute("SELECT name FROM tenants WHERE id=?", (tenant,)).fetchone()
        if row is None:
            return None
        name = (row["name"] or "").strip()
        hotel = con.execute(
            "SELECT id, payload FROM hotels WHERE json_extract(payload, '$.name')=?",
            (name,),
        ).fetchone()
        if hotel is None:
            return None
        import json

        payload = json.loads(hotel["payload"])
        if not (payload.get("coordinates") or {}).get("lng"):
            return None
        return {"id": hotel["id"], "name": payload.get("name"),
                "lng": payload["coordinates"]["lng"], "lat": payload["coordinates"]["lat"]}


def collect(db_path: Path, base_dir: Path, tenant: str,
            today: date | None = None) -> dict:
    """Everything the report shows, gathered before either file format is rendered."""
    start, end = week_range(today)
    with connect(db_path) as con:
        row = con.execute("SELECT name FROM tenants WHERE id=?", (tenant,)).fetchone()
        tenant_name = (row["name"] if row else tenant)
        hotel = _tenant_hotel(db_path, tenant)

        # 去重：同一活动被多个来源收录时只算一场。canonical_id 由 migrate_dedupe 回填，
        # 未迁移的库没有该列时退化为每条记录自成身份。
        columns = {row[1] for row in con.execute("PRAGMA table_info(events)")}
        canonical = "COALESCE(canonical_id, id)" if "canonical_id" in columns else "id"
        events = []
        seen: set[str] = set()
        for event in con.execute(
            f"SELECT id, {canonical} AS event_key, name, category, startDate, endDate, venue, "
            "source, verification_status FROM events WHERE startDate IS NOT NULL"
        ):
            starts = event["startDate"]
            ends = event["endDate"] or starts
            if starts > end.isoformat() or ends < start.isoformat():
                continue
            key = event["event_key"] or event["id"]
            if key in seen:
                continue
            seen.add(key)
            events.append(dict(event, eventKey=key))

        # 本店的跟进状态与本周变化：周报的意义在于「我答应要做什么」和「和上周比变了什么」，
        # 这两块以前完全没有出现在导出文件里。
        followup_rows = []
        try:
            from backend import tracking

            tracking.ensure_schema(con)
            followup_rows = [dict(row) for row in con.execute(
                "SELECT event_key, state, note, updated_at FROM event_followups "
                "WHERE tenant_id=?", (tenant,))]
        except Exception:
            followup_rows = []

        venue_rows = {r["id"]: r["payload"] for r in con.execute("SELECT id, payload FROM venues")}

    change_summary = {"comparable": False, "reason": "尚无快照记录"}
    try:
        from backend import tracking

        change_summary = tracking.changes(db_path)
    except Exception:
        pass

    import json

    venues = {}
    for vid, payload in venue_rows.items():
        venue = json.loads(payload)
        co = venue.get("coordinates") or {}
        venues[vid] = {"name": venue.get("name"), "lng": co.get("lng"), "lat": co.get("lat")}

    scoped = []
    for event in events:
        venue = venues.get(event["venue"] or "", {})
        distance = None
        if hotel and venue.get("lng") is not None:
            distance = round(_haversine_km((hotel["lng"], hotel["lat"]),
                                           (venue["lng"], venue["lat"])), 2)
        scoped.append({
            **event,
            "venueName": venue.get("name"),
            "distanceKm": distance,
        })

    if hotel:
        scoped = [e for e in scoped if e["distanceKm"] is not None and e["distanceKm"] <= NEARBY_KM]
    scoped.sort(key=lambda e: (e["startDate"], e["name"]))

    con = tenants.open_tenant_db(base_dir, tenant)
    try:
        rows = tenants.read_metrics(con)
    finally:
        con.close()

    in_week = [
        r for r in rows
        if r["period_end"] or r["period_start"]
    ]
    in_week = [
        r for r in rows
        if (r["period_start"] <= end.isoformat()
            and (r["period_end"] or r["period_start"]) >= start.isoformat())
    ]

    by_metric: dict[str, list[dict]] = {}
    for row in in_week:
        by_metric.setdefault(row["metric"], []).append(row)

    summary = []
    for metric, items in by_metric.items():
        values = [i["value"] for i in items]
        label, unit = METRIC_LABELS.get(metric, (metric, ""))
        summary.append({
            "metric": metric,
            "label": label,
            "unit": unit,
            "count": len(values),
            "avg": round(sum(values) / len(values), 2),
            "min": min(values),
            "max": max(values),
            "days": sorted({i["period_start"] for i in items}),
        })
    summary.sort(key=lambda s: list(METRIC_LABELS).index(s["metric"])
                 if s["metric"] in METRIC_LABELS else 99)

    hotels_for_name = None
    return {
        "tenant": tenant,
        "tenantName": tenant_name,
        "hotel": hotel,
        "hotelsForName": hotels_for_name,
        "period": {"start": start.isoformat(), "end": end.isoformat(),
                   "label": f"{start.isoformat()} 至 {end.isoformat()}"},
        "summary": summary,
        "rows": in_week,
        "events": scoped,
        "followups": followup_rows,
        "changes": change_summary,
        "todo": [row for row in followup_rows if row.get("state") == "todo"],
        "scope": {
            "nearbyKm": NEARBY_KM if hotel else None,
            "hotelKnown": hotel is not None,
            "note": (
                f"活动范围：距本店 {NEARBY_KM:.0f} km 内的公开活动"
                if hotel
                else "本店位置尚未与场馆库匹配，活动按全市列出；绑定位置后可按距离筛选"
            ),
        },
        "disclaimer": ("本报告只统计本店上传的经营数据与公开活动信息，保留上传人与来源文件；"
                        "不做需求预测、不给房价建议、不估算收入。"),
        "generatedAt": datetime.now(timezone.utc).isoformat(),
    }


# --- Excel -------------------------------------------------------------------------

def build_xlsx(report: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    head = Font(bold=True)

    overview = book.active
    overview.title = "概览"
    overview.append(["本店经营数据周报"])
    overview.cell(row=1, column=1).font = Font(bold=True, size=14)
    overview.append(["酒店", report["tenantName"]])
    overview.append(["统计周期", report["period"]["label"]])
    overview.append(["活动范围", report["scope"]["note"]])
    overview.append(["生成时间", report["generatedAt"][:19].replace("T", " ")])
    overview.append([])
    overview.append(["指标", "记录数", "平均", "最小", "最大", "单位"])
    for cell in overview[7]:
        cell.font = head
    for item in report["summary"]:
        overview.append([item["label"], item["count"], item["avg"], item["min"],
                         item["max"], item["unit"]])
    if not report["summary"]:
        overview.append(["本周没有上传数据"])

    detail = book.create_sheet("经营数据明细")
    detail.append(["指标", "周期起", "周期止", "数值", "渠道", "房型", "上传人", "来源文件"])
    for cell in detail[1]:
        cell.font = head
    for row in report["rows"]:
        detail.append([
            METRIC_LABELS.get(row["metric"], (row["metric"], ""))[0],
            row["period_start"], row["period_end"] or "", row["value"],
            row["channel"] or "", row["room_type"] or "",
            row["uploaded_by"], row["source_file"] or "",
        ])

    events = book.create_sheet("本周活动")
    events.append(["活动", "类别", "开始", "结束", "场馆", "距离(km)", "来源"])
    for cell in events[1]:
        cell.font = head
    for event in report["events"]:
        events.append([
            event["name"], event["category"], event["startDate"],
            event["endDate"] or "", event.get("venueName") or "",
            event.get("distanceKm") if event.get("distanceKm") is not None else "",
            event["source"],
        ])

    changes = book.create_sheet("本周变化")
    changes.append(["类别", "活动", "说明"])
    for cell in changes[1]:
        cell.font = head
    change_summary = report.get("changes") or {}
    if change_summary.get("comparable"):
        for item in change_summary.get("added", []):
            changes.append(["新增", item.get("name"), item.get("start_date") or ""])
        for item in change_summary.get("rescheduled", []):
            before = item.get("before") or {}
            after = item.get("after") or {}
            changes.append([
                "改期", after.get("name"),
                f"{before.get('start_date') or '?'} → {after.get('start_date') or '?'}",
            ])
        for item in change_summary.get("cancelled", []):
            changes.append(["取消", item.get("name"), "快照中已不存在"])
        for item in change_summary.get("changed", []):
            after = item.get("after") or {}
            changes.append(["信息变化", after.get("name"), ""])
        if len(changes["A"]) == 1:
            changes.append(["无变化", "", ""])
    else:
        changes.append(["不可比较", change_summary.get("reason", "尚无快照记录"), ""])

    todo = book.create_sheet("待办")
    todo.append(["活动身份", "状态", "备注", "更新时间"])
    for cell in todo[1]:
        cell.font = head
    for row in report.get("followups") or []:
        todo.append([
            row.get("event_key"), row.get("state"), row.get("note") or "",
            (row.get("updated_at") or "")[:19].replace("T", " "),
        ])
    if not (report.get("followups") or []):
        todo.append(["尚无跟进记录", "", "", ""])

    note = book.create_sheet("说明")
    note.append(["数据范围"])
    note.append([report["scope"]["note"]])
    note.append([])
    note.append([report["disclaimer"]])
    note.append(["来源可追溯：活动记录保留原始来源标识；经营数据保留上传人与来源文件。"])

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


# --- PDF ---------------------------------------------------------------------------

# reportlab only understands TrueType outlines. Noto Sans CJK, the obvious font on Ubuntu,
# is CFF-flavoured and fails with "postscript outlines are not supported", so the list is
# ordered TrueType-first and each candidate is tried rather than assumed.
_FONT_CANDIDATES = (
    "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
    "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
    "C:/Windows/Fonts/msyh.ttc",
    "C:/Windows/Fonts/simsun.ttc",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
)


def register_cjk_font() -> str | None:
    """Register the first candidate reportlab can actually load, and return its name.

    Without a CJK font the PDF would print empty boxes where the Chinese text is, so failure
    is reported instead of shipping an unreadable file.
    """
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for path in _FONT_CANDIDATES:
        if not Path(path).exists():
            continue
        try:
            pdfmetrics.registerFont(TTFont("MendunerCJK", path, subfontIndex=0))
            return "MendunerCJK"
        except Exception:
            continue
    return None


def build_pdf(report: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)

    font_name = register_cjk_font()
    if font_name is None:
        raise RuntimeError("没有 reportlab 可用的中文字体（需要 TrueType 轮廓），无法生成 PDF")

    base = getSampleStyleSheet()
    title = ParagraphStyle("cn-title", parent=base["Title"], fontName=font_name, fontSize=18)
    h2 = ParagraphStyle("cn-h2", parent=base["Heading2"], fontName=font_name, fontSize=13,
                        spaceBefore=10, spaceAfter=4)
    body = ParagraphStyle("cn-body", parent=base["BodyText"], fontName=font_name, fontSize=9.5,
                          leading=14)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                            leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm,
                            title=f"{report['tenantName']} 周报",
                            author="Menduner")
    flow = [Paragraph(f"{report['tenantName']} · 经营数据周报", title),
            Paragraph(f"统计周期：{report['period']['label']}", body),
            Paragraph(f"活动范围：{report['scope']['note']}", body),
            Spacer(1, 6)]

    flow.append(Paragraph("一、本周经营概览", h2))
    if report["summary"]:
        table = [["指标", "记录数", "平均", "最小", "最大", "单位"]]
        for item in report["summary"]:
            table.append([item["label"], str(item["count"]), str(item["avg"]),
                          str(item["min"]), str(item["max"]), item["unit"]])
        t = Table(table, colWidths=[38 * mm, 22 * mm, 26 * mm, 22 * mm, 22 * mm, 18 * mm])
        t.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), font_name),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF2F6")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
            ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ]))
        flow.append(t)
    else:
        flow.append(Paragraph("本周没有上传数据。", body))

    flow.append(Paragraph("二、本周活动", h2))
    if report["events"]:
        table = [["活动", "场馆", "开始", "结束", "距离"]]
        for event in report["events"][:25]:
            table.append([
                Paragraph(event["name"][:34], body),
                Paragraph((event.get("venueName") or "—")[:18], body),
                event["startDate"], event["endDate"] or "—",
                f"{event['distanceKm']} km" if event.get("distanceKm") is not None else "—",
            ])
        t = Table(table, colWidths=[62 * mm, 42 * mm, 22 * mm, 22 * mm, 20 * mm])
        t.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), font_name),
            ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF2F6")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        flow.append(t)
        if len(report["events"]) > 25:
            flow.append(Paragraph(
                f"另有 {len(report['events']) - 25} 个活动，见 Excel 附件。", body))
    else:
        flow.append(Paragraph("本周范围内没有可列出的公开活动。", body))

    flow.append(Paragraph("三、本周变化", h2))
    summary = report.get("changes") or {}
    if summary.get("comparable"):
        lines = []
        for item in summary.get("added", []):
            lines.append(f"新增：{item.get('name')}（{item.get('start_date') or '日期待核验'}）")
        for item in summary.get("rescheduled", []):
            before = item.get("before") or {}
            after = item.get("after") or {}
            lines.append(f"改期：{after.get('name')}（{before.get('start_date') or '?'} → "
                         f"{after.get('start_date') or '?'}）")
        for item in summary.get("cancelled", []):
            lines.append(f"取消：{item.get('name')}")
        for item in summary.get("changed", []):
            lines.append(f"信息变化：{(item.get('after') or {}).get('name')}")
        if not lines:
            lines.append("两份快照之间没有变化。")
        for line in lines[:20]:
            flow.append(Paragraph(line, body))
    else:
        flow.append(Paragraph(summary.get("reason", "尚无快照记录"), body))

    flow.append(Paragraph("四、待办", h2))
    followups = report.get("followups") or []
    if followups:
        for row in followups[:20]:
            flow.append(Paragraph(
                f"{row.get('event_key')} · {row.get('state')}" +
                (f" · {row.get('note')}" if row.get("note") else ""), body))
    else:
        flow.append(Paragraph("本周没有标记跟进的活动。", body))

    flow.append(Spacer(1, 10))
    flow.append(Paragraph("说明：" + report["disclaimer"], body))

    doc.build(flow)
    return buffer.getvalue()
