"""生成周四企业交流用的 PPT 与 Word。

内容来源：docs/ENTERPRISE_MENTOR_BRIEFING.md（GPT Round 11/12 共拟 + DSH 用真实代码与生产库核实）。
设计取向：白底、深灰字、细线分隔，不用彩色块和标签堆叠 —— 汇报现场投影时更干净。
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt as Ppt

OUT = Path(__file__).resolve().parent.parent / "deliverables"
OUT.mkdir(exist_ok=True)

INK = PptRGB(0x1F, 0x29, 0x37)
MUTED = PptRGB(0x6B, 0x72, 0x80)
RULE = PptRGB(0xD8, 0xDE, 0xE6)
ACCENT = PptRGB(0x1D, 0x4E, 0xD8)
FONT = "微软雅黑"

# 真实数据（生产库实测，2026-10-26）
FACTS = {
    "记录": "436",
    "身份": "416",
    "场馆": "135",
    "酒店": "19",
    "来源": "17",
    "证据": "4463",
    "太仓": "16",
    "坐标场馆": "90",
    "快照": "3",
}


# ---------------------------------------------------------------- PPT
def ppt():
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    blank = deck.slide_layouts[6]
    state = {"n": 1}  # 封面已经是第 1 页

    def text(slide, value, left, top, width, height, size, bold=False,
             color=INK, align=PP_ALIGN.LEFT, spacing=1.15):
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        frame = box.text_frame
        frame.word_wrap = True
        for index, line in enumerate(value.split("\n")):
            para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            para.alignment = align
            para.line_spacing = spacing
            run = para.add_run()
            run.text = line
            run.font.size = Ppt(size)
            run.font.bold = bold
            run.font.name = FONT
            run.font.color.rgb = color
        return box

    def rule(slide, top, left=0.9, width=11.53):
        line = slide.shapes.add_connector(1, Inches(left), Inches(top), Inches(left + width), Inches(top))
        line.line.color.rgb = RULE
        line.line.width = Ppt(0.9)

    def page(title, kicker=None):
        state["n"] += 1
        slide = deck.slides.add_slide(blank)
        if kicker:
            text(slide, kicker, 0.9, 0.52, 11.5, 0.32, 12, color=MUTED)
            text(slide, title, 0.9, 0.82, 11.6, 0.7, 30 if len(title) <= 20 else 25, bold=True)
        else:
            text(slide, title, 0.9, 0.62, 11.6, 0.9, 32 if len(title) <= 18 else 27, bold=True)
        rule(slide, 1.62 if kicker else 1.5)
        text(slide, f"{state['n']}", 12.35, 6.95, 0.6, 0.3, 10, color=MUTED, align=PP_ALIGN.RIGHT)
        text(slide, "Menduner · ENT303TC 第 14 组", 0.9, 6.95, 6, 0.3, 10, color=MUTED)
        return slide

    def bullets(slide, items, top=1.95, size=17, gap=0.62, left=0.95, width=11.3):
        for index, item in enumerate(items):
            mark, body = ("· ", item) if not item.startswith("—") else ("", item[1:].strip())
            y = top + index * gap
            if mark:
                text(slide, mark, left, y, 0.3, 0.4, size, color=ACCENT)
            text(slide, body, left + 0.34, y, width - 0.34, 0.5, size)

    def table(slide, header, rows, top=2.0, widths=None, size=13, row_h=0.46):
        """无填充、无网格的表格：只用细线分隔，避免默认蓝色表头（投影时更干净）。"""
        from pptx.oxml.ns import qn as pqn

        cols = len(header)
        total = 11.3
        widths = widths or [total / cols] * cols
        shape = slide.shapes.add_table(len(rows) + 1, cols, Inches(0.95), Inches(top),
                                       Inches(sum(widths)), Inches(row_h * (len(rows) + 1)))
        grid = shape.table
        # 移除默认表格样式（蓝色表头 + 斑马纹）
        tbl = grid._tbl
        tblPr = tbl.tblPr
        for attr in ("firstRow", "bandRow", "firstCol"):
            if tblPr.get(attr) is not None:
                tblPr.set(attr, "0")
        style_el = tblPr.find(pqn("a:tableStyleId"))
        if style_el is not None:
            style_el.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
        else:
            from lxml import etree

            node = etree.SubElement(tblPr, pqn("a:tableStyleId"))
            node.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"

        for c, w in enumerate(widths):
            grid.columns[c].width = Inches(w)
        for r in range(len(rows) + 1):
            grid.rows[r].height = Inches(row_h)

        def fill(cell, value, bold=False, color=INK, cell_size=None):
            cell.text = ""
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            para = cell.text_frame.paragraphs[0]
            para.line_spacing = 1.05
            run = para.add_run()
            run.text = str(value)
            run.font.size = Ppt((cell_size or size) - 1.5)
            run.font.bold = bold
            run.font.name = FONT
            run.font.color.rgb = color

        for c, label in enumerate(header):
            fill(grid.cell(0, c), label, bold=True, color=MUTED, cell_size=size - 1.5)
        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row):
                fill(grid.cell(r, c), value, color=INK)

        # 表头下的细线
        bottom = top + row_h
        line = slide.shapes.add_connector(1, Inches(0.95), Inches(bottom),
                                          Inches(0.95 + sum(widths)), Inches(bottom))
        line.line.color.rgb = RULE
        line.line.width = Ppt(0.9)
        return grid

    # 1 封面
    slide = deck.slides.add_slide(blank)
    text(slide, "Menduner", 0.95, 2.05, 8, 0.9, 44, bold=True)
    text(slide, "把城市公开活动，整理成酒店能用的跟进清单", 0.95, 2.95, 11.4, 0.6, 22, color=INK)
    rule(slide, 3.85, 0.95, 5.2)
    text(slide, "周四企业导师交流 · 成果汇报与需求验证", 0.95, 4.05, 11, 0.45, 15, color=MUTED)
    text(slide, "ENT303TC 第 14 组 · 2026 年 10 月", 0.95, 4.5, 11, 0.45, 15, color=MUTED)
    text(slide, "在线版本：entgroup14.com", 0.95, 5.65, 11, 0.4, 13, color=MUTED)

    # 2 今天的目的
    slide = page("今天我们主要不是来展示软件", "交流目标")
    table(slide, ["时间分配", "我们想做的事"], [
        ["30%", "汇报已经做出来的东西（只说真实上线的）"],
        ["40%", "请您讲酒店实际怎么获取和使用活动信息"],
        ["30%", "争取一个成本低、范围小的验证机会"],
    ], top=2.05, widths=[2.2, 9.1], size=13, row_h=0.56)
    text(slide, "我们不谈采购，也不承诺提升收入。先把方向做对。", 0.95, 4.9, 11.3, 0.5, 16, color=MUTED)

    # 3 一句话 + 数字
    slide = page("Menduner 是什么：一个可追溯的城市活动工作流", "产品定位")
    bullets(slide, [
        "采集城市公开活动 → 去重 → 关联场馆与地区 → 判断对哪家酒店有用 → 形成跟进清单与周报",
        "每个数字都能点回来源；说不清的字段标「待补充」，不猜",
        "面向酒店市场/销售/收益岗位的日常信息整理，不替代任何收益系统",
    ], top=1.95, size=17, gap=0.66)
    stats = [("活动记录", FACTS["记录"]), ("去重后身份", FACTS["身份"]),
             ("场馆", FACTS["场馆"]), ("酒店", FACTS["酒店"]),
             ("公开来源", FACTS["来源"]), ("字段证据", FACTS["证据"])]
    for i, (label, value) in enumerate(stats):
        x = 0.95 + i * 1.9
        text(slide, value, x, 4.55, 1.8, 0.6, 30, bold=True)
        text(slide, label, x, 5.18, 1.8, 0.35, 12, color=MUTED)
    text(slide, "数据为 2026-10-26 生产环境实测；太仓市直接归属活动 16 条 —— 这是样本，不是全覆盖。",
         0.95, 5.9, 11.3, 0.4, 13, color=MUTED)

    # 4 数据从哪来
    slide = page("数据为什么可信：来源、证据、状态三件事", "数据基础")
    table(slide, ["我们怎么做的", "对酒店意味着什么"], [
        ["17 个公开来源（政府、场馆、活动平台、旅游咨询）；聚合站与转载不算独立来源",
         "一条活动至少能追到第一条公开出处"],
        ["4463 条字段证据可点开；分关键/辅助两组，关键字段缺失标待核验", "既看到结论与依据，也知道哪条还需确认"],
        [f"每日快照（当前 {FACTS['快照']} 份），同一天的数据可比", "能看到活动是新增、改期还是消失"],
    ], top=2.0, widths=[7.3, 4.0], size=12.5, row_h=0.8)
    text(slide, "字段已核验 ≠ 整条活动已核验 —— 这句话我们写在了页面上。", 0.95, 6.15, 11.3, 0.4, 14, color=MUTED)

    # 5 演示路径
    slide = page("接下来 13 分钟：只看三个动作", "演示路径")
    steps = [("① 发现活动", "地图 / 清单：这家酒店周边有什么、什么时候"),
             ("② 酒店跟进", "工作台：哪些值得跟进、谁在跟进、和上次比变了什么"),
             ("③ 导出简报", "商圈洞察：一份可进周会的 Excel，每行带来源")]
    for i, (head, body) in enumerate(steps):
        y = 2.1 + i * 1.35
        text(slide, head, 1.0, y, 3.4, 0.5, 22, bold=True)
        text(slide, body, 4.5, y + 0.06, 7.7, 0.5, 16, color=MUTED)
        rule(slide, y + 0.95, 1.0, 11.2)
    text(slide, "全程只用一个城市（太仓）+ 一家真实收录酒店 + 一条真实活动 + 一份导出文件。",
         0.95, 6.35, 11.3, 0.4, 14, color=MUTED)

    # 6 动作一
    slide = page("① 发现活动：地图与清单", "演示 1/3")
    bullets(slide, [
        "按年份、类别、区域筛选；演示数据默认关闭，已结束活动默认隐藏",
        "月初默认只看本月，月份控件折叠起来不占位置（1440×900 下地图可见高度 44%）",
        "清单按「即将开始 / 进行中」分组；缺坐标的活动单独标注，不从统计里消失",
        "每个活动点开能看到来源、核验状态与缺失字段",
    ], top=1.95, size=17, gap=0.72)
    text(slide, "关键句：我们整理的是公开活动信息，不是城市全部活动。", 0.95, 5.55, 11.3, 0.5, 16, color=ACCENT)

    # 7 动作二
    slide = page("② 酒店跟进：把这些活动变成待办", "演示 2/3")
    bullets(slide, [
        "工作台给每家店列出未来 7 天值得关注的活动，并写明推荐理由",
        "四个状态：关注 / 待跟进 / 已跟进 / 忽略；另有「是否采纳建议」记录",
        "推荐用四维规则（业务相关性、距离、数据可信度、紧迫度）；缺数据的维度显示「未知」",
        "同一家店的员工看到的是同一份状态，不用在群里反复对口径",
    ], top=1.95, size=17, gap=0.72)

    # 8 动作三
    slide = page("③ 导出简报：一份能进周会的 Excel", "演示 3/3")
    table(slide, ["工作表", "内容"], [
        ["城市概览", "地区、时间段、去重后活动数、覆盖天数、来源数、批次"],
        ["活动明细", "名称、类别、起止、场馆、地区、坐标有无、来源——每行可回溯源"],
        ["本期变化", "只在有两份可比快照时给出，否则写明暂不提供及原因"],
        ["证据与待确认", "来源名称、公开链接、核对时间；缺场馆/坐标/日期的单列一行"],
    ], top=1.95, widths=[2.3, 9.0], size=12, row_h=0.52)

    # 9 其余页面
    slide = page("其他功能：老师想深入看时我们再打开", "附录")
    table(slide, ["页面", "回答什么问题"], [
        ["知识图谱", "一个活动和哪些场馆、地区、酒店有关；缺的关系写「待补充」"],
        ["商圈洞察", "行政区与商业商圈分开统计；经济情景是参数假设，缺参数不出金额"],
        ["酒店活动情报", "某家酒店周边 10 km 内的活动；坐标未核验的不计入距离"],
        ["找酒店", "按星级、地区、名称找已收录的 19 家酒店"],
        ["本店经营数据", "酒店自己上传经营数据，只在本店账号内可见"],
        ["数据质量与来源", "每条数据能不能用、来源在哪、还缺什么"],
    ], top=1.95, widths=[2.5, 8.8], size=12, row_h=0.5)

    # 10 边界
    slide = page("我们明确不做的事", "边界")
    table(slide, ["不做", "原因"], [
        ["需求预测、房价建议", "没有足够的酒店侧数据，做了也不可信"],
        ["替代酒店收益管理系统（RMS）", "我们做的是它的输入侧，不是竞品"],
        ["承诺提升 RevPAR / ADR", "现在没有证据，承诺就是不专业"],
        ["把参数情景说成真实收入", "假设和观测必须分开说"],
    ], top=2.0, widths=[4.6, 6.7], size=12.5, row_h=0.66)
    text(slide, "能把「不做什么」说清楚，比多讲十个功能更可信。", 0.95, 5.6, 11.3, 0.5, 16, color=MUTED)

    # 11 想请老师帮的三件事
    slide = page("想请老师帮的三件事", "合作请求")
    table(slide, ["层级 · 请求", "我们的承诺"], [
        ["第一层：帮我们评审一份真实城市简报，指出无效字段与缺失信息", "不增加企业同事工作量"],
        ["第二层：20–30 分钟，与市场/销售/收益岗位交流业务流程", "提前给议题，不超时"],
        ["第三层：推荐一家合适酒店，做单城市小范围验证", "系统运行、简报整理、问题修正由我们负责"],
    ], top=1.95, widths=[6.6, 4.7], size=13, row_h=0.78)
    text(slide, "不需要接入酒店内部系统，也不要客户名单、订单明细或任何个人信息。",
         0.95, 5.35, 11.3, 0.45, 15, color=ACCENT)
    text(slide, "如果近期不方便，也完全可以先从业务建议和材料反馈开始。", 0.95, 5.85, 11.3, 0.4, 14, color=MUTED)

    # 12 企业能得到什么
    slide = page("企业能得到的：三项可以量出来的结果", "价值")
    table(slide, ["维度", "验证方式"], [
        ["信息价值", "能否更早发现与自身相关的活动（发现时间差）"],
        ["工作效率", "整理一份活动简报所需时间是否减少"],
        ["决策价值", "员工能否基于简报挑出该跟进的活动，并说清理由"],
    ], top=2.0, widths=[2.6, 8.7], size=13, row_h=0.6)
    text(slide, "我们不承诺营业额提升，只承诺把这三件事量出来给您看。", 0.95, 4.75, 11.3, 0.5, 16, color=MUTED)

    # 13 路线图
    slide = page("后续路线：三步走", "路线图")
    stages = [
        ("短期 2–4 周 · 收尾", "交互测试收尾、使用说明、演示流程、数据覆盖说明", "交付：使用说明书 + 试点需求清单"),
        ("中期 1–3 月 · 试点", "按约定范围整理活动、生成简报、记录问题、改进质量", "交付：试点评估报告 + 功能优先级"),
        ("长期 · 商业化探索", "数据维护成本、标准化、多店管理、权限隔离", "交付：基于真实使用反馈的可行性方案"),
    ]
    for index, (head, body_text, deliver) in enumerate(stages):
        y = 1.9 + index * 1.62
        text(slide, head, 0.95, y, 11.3, 0.42, 20, bold=True)
        text(slide, body_text, 0.95, y + 0.48, 11.3, 0.4, 15, color=INK)
        text(slide, deliver, 0.95, y + 0.86, 11.3, 0.36, 13, color=MUTED)
        rule(slide, y + 1.28, 0.95, 11.2)
    text(slide, "每一步都需要企业配合：短期给业务反馈，中期给一个试点场景，长期给采购与预算流程的指导。",
         0.95, 6.35, 11.3, 0.4, 13, color=MUTED)

    # 14 离场想拿到的
    slide = page("今天离场前，希望能拿到两个答案", "收尾")
    bullets(slide, [
        "先找哪个岗位、先解决哪一个具体问题？",
        "下一步由谁在什么时候给我们一次反馈（哪怕是微信一句话）？",
    ], top=2.2, size=21, gap=0.95)
    text(slide, "如果只能给一个建议：您觉得我们最该先验证的方向是哪个？", 0.95, 4.75, 11.3, 0.5, 17, color=ACCENT)

    # 15 谢谢
    slide = deck.slides.add_slide(blank)
    text(slide, "谢谢老师", 0.95, 2.45, 8, 0.9, 40, bold=True)
    text(slide, "系统在线体验：entgroup14.com", 0.95, 3.5, 11, 0.45, 17, color=MUTED)
    text(slide, "后续材料与试点方案，我们会在一周内整理后发您。", 0.95, 4.0, 11, 0.45, 17, color=MUTED)
    rule(slide, 4.75, 0.95, 5.2)

    path = OUT / "Menduner_企业导师交流.pptx"
    deck.save(path)
    return path


# ---------------------------------------------------------------- DOCX
def docx_file():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(2.0)
        section.left_margin = section.right_margin = Cm(2.2)

    def heading(value, level=1):
        para = doc.add_heading(value, level=level)
        for run in para.runs:
            run.font.name = "微软雅黑"
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        para.paragraph_format.space_before = Pt(10)
        para.paragraph_format.space_after = Pt(4)
        return para

    def body(value, size=10.5, bold=False, color=None, space=6):
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(space)
        para.paragraph_format.line_spacing = 1.35
        run = para.add_run(value)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        if color:
            run.font.color.rgb = color
        return para

    def table(header, rows, widths=None, size=9.5):
        grid = doc.add_table(rows=len(rows) + 1, cols=len(header))
        grid.style = "Table Grid"
        grid.alignment = WD_TABLE_ALIGNMENT.CENTER
        grid.autofit = False
        # 固定列宽：auto 布局会在中文字体下把表格撑出页面
        from docx.oxml import OxmlElement

        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        grid._tbl.tblPr.append(layout)
        for c, label in enumerate(header):
            cell = grid.cell(0, c)
            cell.text = ""
            run = cell.paragraphs[0].add_run(label)
            run.font.bold = True
            run.font.size = Pt(size)
            run.font.name = "微软雅黑"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row):
                cell = grid.cell(r, c)
                cell.text = ""
                run = cell.paragraphs[0].add_run(str(value))
                run.font.size = Pt(size)
                run.font.name = "宋体"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        if widths:
            for c, width in enumerate(widths):
                for row in grid.rows:
                    row.cells[c].width = Cm(width)
        doc.add_paragraph()
        return grid

    title = doc.add_heading("Menduner 企业导师交流：问题清单与试点合作说明", level=0)
    for run in title.runs:
        run.font.name = "微软雅黑"
        run.font.size = Pt(17)
        run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    body("ENT303TC 第 14 组 · 2026 年 10 月 · 在线版本 entgroup14.com", size=9, color=RGBColor(0x6B, 0x72, 0x80))
    body("本文件用于周四交流现场提问与后续沟通留档。所有数字来自生产环境实测（2026-10-26），"
         "已上线功能与未做/待验证内容严格区分，不做夸大表述。")

    heading("一、真实数据与能力边界（沟通时的统一口径）", 1)
    table(["项目", "真实值", "说明"], [
        ["活动记录", "436 条", "公开来源采集"],
        ["去重后活动身份", "416 个", "跨来源同一活动合并计数"],
        ["场馆 / 酒店 / 地区", "135 / 19 / 39", "场馆有坐标 90 个"],
        ["公开来源", "17 个", "政府、场馆、活动平台、旅游咨询等"],
        ["字段证据", "4463 条", "逐字段可追到原链接"],
        ["太仓市直接归属活动", "16 条", "样本量有限，不作为覆盖证明"],
        ["每日快照", "3 份", "支持「本期变化」比对"],
    ], widths=[4.2, 3.2, 8.4])
    body("对外口径：我们整理的是可追溯的公开活动信息，不是城市全部活动的覆盖量。", bold=True)
    table(["状态", "包含内容"], [
        ["已上线", "活动地图、活动清单、本店工作台、知识图谱、商圈洞察、数据质量与来源、城市简报 Excel 导出"],
        ["待验证", "推荐排序的实际业务价值、简报字段是否满足酒店周会需要、经济情景参数的取值来源"],
        ["未做", "完整的一周变化周报、自动核验流水线、机场类场所数据、公开字段全量白名单"],
    ], widths=[2.6, 13.2])

    heading("二、给企业老师的 10 个问题（现场挑 5–6 个）", 1)
    body("顺序原则：先请教实际工作 → 再谈产品价值 → 最后自然过渡到试点。", size=9.5,
         color=RGBColor(0x6B, 0x72, 0x80))
    table(["#", "问题", "我们想获得的信息"], [
        ["1", "酒店目前通过哪些渠道了解周边展会、赛事和城市活动？", "现有信息获取流程"],
        ["2", "这些信息一般由市场部、销售部、收益管理部还是前厅部负责？", "实际使用者与决策者"],
        ["3", "有没有出现过活动信息发现太晚，导致错过营销或经营准备的情况？", "真实业务痛点"],
        ["4", "更关心活动时间与位置，还是活动规模、客群和潜在客房需求？", "核心功能优先级"],
        ["5", "如果每周发送一份城市活动简报，哪些字段是必须包含的？", "周报的业务标准"],
        ["6", "什么样的来源证据才足以让酒店采用一条活动信息？", "数据可信度门槛"],
        ["7", "假设系统推荐了十场活动，酒店会依据什么决定重点跟进其中两三场？", "推荐与筛选逻辑"],
        ["8", "酒店是否会评估大型活动对入住率、ADR 或 RevPAR 的影响？通常如何评估？", "经济分析的现实基础"],
        ["9", "如果做小规模试点，您认为应该从哪个部门、哪类酒店或哪个场景开始？", "试点切入点"],
        ["10", "如果试点有效，企业内部通常还需要经过哪些环节才能正式采用？", "采购、合规、预算流程"],
    ], widths=[1.0, 9.6, 5.2])

    heading("三、现场话术（可直接照念）", 1)
    body("开场：", bold=True, space=2)
    body("“老师，我们能采集和整理活动，但还不确定酒店工作中最需要哪一部分。想先请教您平时怎么处理这些信息，"
         "避免我们只按技术思路继续开发。”")
    body("老师指出不足时：", bold=True, space=2)
    body("“这个问题很重要。我们之前更多从数据和技术角度考虑，实际业务流程确实需要向企业学习。我们会单独记录下来，"
         "再判断是调整功能，还是先通过试点验证。”")
    body("最可能被质疑的一点（“系统不错，但酒店不一定需要”）：", bold=True, space=2)
    body("“我们这次最想验证的不是功能有多少，而是它能不能帮酒店减少信息整理时间、提前发现值得跟进的活动；"
         "您觉得哪个岗位最值得先试用？”")
    body("避免说的：", bold=True, space=2)
    body("“我们技术上都能实现。” —— 企业导师更关心你们是否理解实际工作场景与交付约束。")

    heading("四、试点资源请求（由轻到重，四级）", 1)
    table(["层级", "请求内容", "我们的承诺"], [
        ["第一层", "评审一份真实城市简报，指出无效字段与缺失信息", "不增加企业同事工作量"],
        ["第二层", "20–30 分钟与相关岗位交流业务流程", "提前给议题，不超时"],
        ["第三层", "推荐一家合适酒店，做单城市或酒店周边范围验证", "系统运行、简报整理、问题修正由我们负责"],
        ["第四层", "试点评价指标定义与公开或汇总数据", "不要客户名单、订单明细、个人信息"],
    ], widths=[2.0, 8.4, 5.4])
    body("可直接使用的一段话：", bold=True, space=2)
    body("“老师，我们想再请教一个后续落地的问题。目前系统已经能完成城市活动整理、来源追溯和简报导出，"
         "但这些功能究竟能否真正帮酒店提高效率，还需要实际业务场景来验证。如果企业这边方便，我们希望先请一位熟悉"
         "酒店市场、销售或者收益管理工作的老师，帮我们看一份真实的城市活动简报。如果反馈比较积极，再考虑选择一家"
         "合适的酒店做小范围试点。初期不需要接入酒店内部系统，也不需要提供客户信息或者敏感经营数据，我们主要想了解"
         "实际使用流程和存在的问题。系统运行、简报整理、问题修正以及试点结果汇报由我们负责，尽量不增加企业同事的"
         "额外工作量。当然，如果近期不方便安排试点，也完全可以先从业务建议和材料反馈开始。”")
    body("企业可以获得什么（三项可量化的结果）：", bold=True, space=2)
    table(["维度", "验证方式"], [
        ["信息价值", "能否更早发现与自身相关的活动（发现时间差）"],
        ["工作效率", "整理一份活动简报所需时间是否减少"],
        ["决策价值", "员工能否基于简报挑出该跟进的活动，并说清理由"],
    ], widths=[3.0, 12.8])
    body("收尾可以只问一句：“如果只能推荐一个最适合我们先验证的方向，您会建议找酒店的哪个岗位、解决哪一个具体问题？”")

    heading("五、后续路线图", 1)
    table(["阶段", "我们负责", "希望企业配合", "交付成果"], [
        ["短期 2–4 周", "关键交互测试收尾、逐按钮使用说明、演示流程、数据覆盖说明",
         "业务反馈、岗位工作流程、简报字段建议", "产品使用说明书 + 试点需求清单 + 可演示版本"],
        ["中期 1–3 个月", "按约定范围整理活动、生成简报、记录问题、改进数据质量与体验",
         "一个试点场景、一位业务反馈联系人、必要的非敏感指标", "试点评估报告 + 功能优先级 + 数据质量改善方案"],
        ["长期（试点有效后）", "数据维护成本、产品标准化、多酒店管理、权限隔离、商业模式",
         "对采购流程、预算主体、使用场景与行业推广方式的指导", "基于真实使用反馈的商业化可行性方案"],
    ], widths=[2.6, 5.0, 4.4, 3.8])

    heading("六、常见追问的最低标准答案", 1)
    table(["追问", "回答"], [
        ["数据从哪来？可信吗？", "17 个公开来源（政府、场馆、活动平台、旅游咨询），4463 条字段证据，每条可点回原链接；聚合站与转载不算独立来源。"],
        ["覆盖多少活动？", "436 条记录、416 个身份；太仓直接归属 16 条。这是样本，不是全覆盖。"],
        ["酒店数据怎么拿到？", "现在只有 19 家酒店的公开信息；经营数据需酒店自己上传，只在本店账号内可见。"],
        ["和现有 RMS / 收益系统什么关系？", "不替代。我们做活动信息、来源追溯与简报，是收益系统的输入侧。"],
        ["推荐凭什么可信？", "四维规则（业务相关性、距离、数据可信度、紧迫度），权重具名可查；缺数据标未知，不补值。"],
        ["能提升多少收入？", "现在不能回答，也不该回答。要先把「是否被使用、节省多少时间」用试点量出来。"],
    ], widths=[4.2, 11.6])

    heading("七、明确不做的四件事", 1)
    table(["不做", "原因"], [
        ["需求预测、房价建议", "缺少足够的酒店侧数据，做了也不可信"],
        ["替代酒店收益管理系统（RMS）", "我们做的是输入侧，不是竞品"],
        ["承诺提升 RevPAR / ADR", "没有证据的承诺不专业"],
        ["把参数情景说成真实收入", "假设与观测必须分开表述"],
    ], widths=[5.2, 10.6])

    body("本文件中「已上线」功能均可在 entgroup14.com 现场验证；「待验证」与「未做」项请如实告知企业老师。",
         size=9, color=RGBColor(0x6B, 0x72, 0x80))

    path = OUT / "Menduner_企业导师交流_问题与试点说明.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    for result in (ppt(), docx_file()):
        print("  已生成:", result, f"({result.stat().st_size // 1024} KB)")
